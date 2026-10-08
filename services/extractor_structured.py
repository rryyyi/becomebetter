"""Word/Excel/结构化导出/压缩包解析器，输出统一 Document IR。"""
import io
import zipfile
from pathlib import Path


def extract_docx(path: str) -> dict:
    from docx import Document
    doc = Document(path)
    blocks = [{"type": "text", "page": 1, "text": p.text.strip(), "bbox": None}
              for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            blocks.append({"type": "table", "page": 1, "text": "\n".join(rows), "bbox": None})
    return {"blocks": blocks, "images": [], "metadata": {"format": "docx", "page_count": 1}}


def extract_xlsx(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                rows.append(" | ".join(values))
        if rows:
            blocks.append({"type": "table", "sheet": ws.title, "page": 1,
                           "text": "\n".join(rows), "bbox": None})
    return {"blocks": blocks, "images": [], "metadata": {"format": "xlsx", "sheets": wb.sheetnames}}


def extract_archive(path: str) -> dict:
    blocks = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith("/") or Path(name).name.startswith("."):
                continue
            suffix = Path(name).suffix.lower()
            raw = archive.read(name)
            if suffix in {".txt", ".csv", ".xml", ".json", ".eplan", ".tia"}:
                text = raw.decode("utf-8", errors="ignore").strip()
                if text:
                    blocks.append({"type": "text", "source_name": name, "page": 1,
                                   "text": text, "bbox": None})
    return {"blocks": blocks, "images": [], "metadata": {"format": "archive", "file_count": len(blocks)}}
